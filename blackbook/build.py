"""Build the Folium blackbook: Markdown chapters -> styled .docx -> .pdf.

Usage (from the repository root):
    python blackbook/build.py

Needs pandoc, python-docx and LibreOffice (soffice) on PATH. The .docx is the
editable master; the PDF is rendered by LibreOffice after refreshing the table
of contents and the lists of figures and tables.
"""
import copy
import glob
import os
import re
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Mm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
CHAPTERS = sorted(glob.glob(os.path.join(HERE, "chapters", "*.md")))
OUT_DOCX = os.path.join(HERE, "Folium_Blackbook.docx")
OUT_PDF = os.path.join(HERE, "Folium_Blackbook.pdf")
REF_DOCX = os.path.join(HERE, "reference.docx")
SERIF = "Times New Roman"
MONO = "Consolas"
TAB_RIGHT = 8640  # twips, right edge of the text block

# --------------------------------------------------------------------------- #
# 1. Reference document (styles, page set-up, footer)
# --------------------------------------------------------------------------- #

def set_font(style, size=None, bold=None, italic=None, name=SERIF, color=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    f.color.rgb = color or RGBColor(0, 0, 0)


def para(style, align=None, line=None, before=None, after=None, keep_next=None, page_break=None,
         indent_first=None):
    pf = style.paragraph_format
    if align is not None:
        pf.alignment = align
    if line is not None:
        pf.line_spacing = line
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if page_break is not None:
        pf.page_break_before = page_break
    if indent_first is not None:
        pf.first_line_indent = indent_first


def get_style(doc, name, kind=1):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, kind)


def add_page_field(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def build_reference():
    subprocess.run(["pandoc", "-o", REF_DOCX, "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(REF_DOCX)
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin, s.right_margin = Inches(1.25), Inches(1.0)
    s.top_margin, s.bottom_margin = Inches(1.0), Inches(1.0)
    s.footer_distance = Inches(0.5)
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)
    for r in fp.runs:
        r.font.name, r.font.size = SERIF, Pt(11)

    body = (WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5)
    for name in ("Normal", "Body Text", "First Paragraph"):
        st = get_style(doc, name)
        set_font(st, 12)
        para(st, align=body[0], line=body[1], before=0, after=6)
    for name, size, before, after, brk in (("Heading 1", 16, 0, 18, True), ("Heading 2", 14, 12, 6, False),
                                           ("Heading 3", 12, 10, 4, False)):
        st = get_style(doc, name)
        set_font(st, size, bold=True, italic=False)
        para(st, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.15, before=before, after=after, keep_next=True,
             page_break=brk)
    for name in ("Image Caption", "Table Caption", "Listing Caption"):
        st = get_style(doc, name)
        set_font(st, 11, bold=False, italic=True)
        para(st, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.15, before=4, after=10,
             keep_next=(name != "Image Caption"))
    for name in ("Captioned Figure", "Figure", "Cover Logo"):
        st = get_style(doc, name)
        para(st, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0, before=6, after=2, keep_next=True)
    st = get_style(doc, "Compact")
    set_font(st, 10.5)
    para(st, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.1, before=1, after=1)
    st = get_style(doc, "Source Code")
    set_font(st, 8.5, name=MONO)
    para(st, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, before=0, after=8)
    st = get_style(doc, "Cover Title")
    set_font(st, 17, bold=True)
    para(st, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.3, before=24, after=24)
    st = get_style(doc, "Cover Text")
    set_font(st, 13)
    para(st, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.15, before=0, after=4)
    st = get_style(doc, "Front Heading")
    set_font(st, 16, bold=True)
    para(st, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.15, before=12, after=24, keep_next=True)
    st = get_style(doc, "Signature Line")
    set_font(st, 12)
    para(st, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.15, before=0, after=0)
    for name in ("TOC 1", "TOC 2"):
        st = get_style(doc, name)
        set_font(st, 12 if name == "TOC 1" else 11.5, bold=False)
        para(st, line=1.25, before=(6 if name == "TOC 1" else 0), after=0)
    doc.save(REF_DOCX)


# --------------------------------------------------------------------------- #
# 2. Markdown pre-processing (markers -> raw OpenXML)
# --------------------------------------------------------------------------- #

def raw(xml):
    return "\n```{=openxml}\n" + xml + "\n```\n"


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def runs(text):
    out = []
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        b = "<w:b/>" if i % 2 else ""
        out.append(f'<w:r><w:rPr>{b}</w:rPr><w:t xml:space="preserve">{esc(part)}</w:t></w:r>')
    return "".join(out)


def sig_line(left, right=None, align="left"):
    jc = {"left": "left", "center": "center", "right": "right"}[align]
    tabs = f'<w:tabs><w:tab w:val="right" w:pos="{TAB_RIGHT}"/></w:tabs>' if right is not None else ""
    body = runs(left) + ("<w:r><w:tab/></w:r>" + runs(right) if right is not None else "")
    return (f'<w:p><w:pPr><w:pStyle w:val="SignatureLine"/>{tabs}<w:jc w:val="{jc}"/></w:pPr>{body}</w:p>')


def field_list(title, instr):
    return raw(
        f'<w:p><w:pPr><w:pStyle w:val="FrontHeading"/><w:pageBreakBefore/></w:pPr>'
        f'<w:r><w:t>{title}</w:t></w:r></w:p>'
        '<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>[Right-click and choose Update Field to refresh this list.]</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')


def preprocess(text):
    out = []
    for line in text.split("\n"):
        m = re.fullmatch(r"\[\[(\w+)(?::(.*))?\]\]", line.strip())
        if not m:
            out.append(line)
            continue
        kind, arg = m.group(1), m.group(2)
        if kind == "PAGEBREAK":
            out.append(raw('<w:p><w:r><w:br w:type="page"/></w:r></w:p>'))
        elif kind == "BLANK":
            out.append(raw('<w:p><w:pPr><w:pStyle w:val="SignatureLine"/></w:pPr></w:p>' * int(arg)))
        elif kind == "SIG":
            left, right = arg.split("|")
            out.append(raw(sig_line(left, right)))
        elif kind in ("SIGL", "SIGC", "SIGR"):
            out.append(raw(sig_line(arg, None, {"SIGL": "left", "SIGC": "center", "SIGR": "right"}[kind])))
        elif kind == "SECTION":
            out.append(raw(f'<w:p><w:r><w:t>@@SECTION_END_{"TITLE" if arg == "roman" else "ROMAN"}@@</w:t></w:r></w:p>'))
        elif kind == "TOC":
            out.append(field_list("Table of Contents", r'TOC \h \z \t "Heading 1,1,Heading 2,2"'))
        elif kind == "LOF":
            out.append(field_list("List of Figures", r'TOC \h \z \t "Image Caption,1"'))
        elif kind == "LOT":
            out.append(field_list("List of Tables", r'TOC \h \z \t "Table Caption,1"'))
        else:
            raise ValueError(f"unknown marker {line}")
    return "\n".join(out)


# --------------------------------------------------------------------------- #
# 3. Post-processing of the generated .docx
# --------------------------------------------------------------------------- #

def set_cell_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    tblPr.append(borders)
    w = tblPr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW")
        tblPr.append(w)
    w.set(qn("w:type"), "pct")
    w.set(qn("w:w"), "5000")
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblPr.append(jc)
    # header row shading and repeat
    if tbl.rows:
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        trPr.append(th)
        for cell in tbl.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "E3F1EF")
            tcPr.append(shd)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True


def shade_code(doc):
    for p in doc.paragraphs:
        if p.style is not None and p.style.name == "Source Code":
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F4F6F6")
            pPr.append(shd)
            bdr = OxmlElement("w:pBdr")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:space"), "4")
                el.set(qn("w:color"), "BFC7C9")
                bdr.append(el)
            pPr.append(bdr)


def sections(doc):
    body = doc.element.body
    final = body.find(qn("w:sectPr"))
    pg = final.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        final.append(pg)
    pg.set(qn("w:fmt"), "decimal")
    pg.set(qn("w:start"), "1")
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt not in ("@@SECTION_END_TITLE@@", "@@SECTION_END_ROMAN@@"):
            continue
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        sp = copy.deepcopy(final)
        for t in sp.findall(qn("w:type")):
            sp.remove(t)
        typ = OxmlElement("w:type")
        typ.set(qn("w:val"), "nextPage")
        sp.insert(0, typ)
        pgn = sp.find(qn("w:pgNumType"))
        if txt == "@@SECTION_END_TITLE@@":
            for ref in sp.findall(qn("w:footerReference")) + sp.findall(qn("w:headerReference")):
                sp.remove(ref)
            sp.remove(pgn)
        else:
            pgn.set(qn("w:fmt"), "lowerRoman")
            pgn.set(qn("w:start"), "1")
        pPr = p._p.get_or_add_pPr()
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "20")
        spacing.set(qn("w:lineRule"), "exact")
        pPr.append(spacing)
        pPr.append(sp)


def update_fields_on_open(doc):
    settings = doc.settings.element
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


def list_paragraphs_to_body(doc):
    """Tight lists use 'Compact' (meant for table cells); give them body text style."""
    body = doc.styles["Body Text"]
    for p in doc.paragraphs:
        if p.style is not None and p.style.name == "Compact":
            p.style = body
            p.paragraph_format.space_after = Pt(3)


def postprocess(path):
    doc = Document(path)
    list_paragraphs_to_body(doc)
    for t in doc.tables:
        set_cell_borders(t)
    shade_code(doc)
    sections(doc)
    update_fields_on_open(doc)
    doc.save(path)


# --------------------------------------------------------------------------- #
# 4. Render the PDF with LibreOffice (indexes refreshed)
# --------------------------------------------------------------------------- #

UNO_SCRIPT = r'''
import sys, time, subprocess, uno
from com.sun.star.beans import PropertyValue
def prop(n, v):
    p = PropertyValue(); p.Name = n; p.Value = v; return p
src, pdf = sys.argv[1], sys.argv[2]
proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore",
    "--accept=socket,host=localhost,port=2002;urp;"])
ctx = None
for _ in range(60):
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        ctx = resolver.resolve("uno:socket,host=localhost,port=2002;urp;StarOffice.ComponentContext")
        break
    except Exception:
        time.sleep(1)
smgr = ctx.ServiceManager
desktop = smgr.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(src), "_blank", 0, (prop("Hidden", True),))
for _ in range(2):
    idx = doc.getDocumentIndexes()
    for i in range(idx.getCount()):
        idx.getByIndex(i).update()
    doc.refresh()
print("indexes", doc.getDocumentIndexes().getCount())
doc.storeToURL(uno.systemPathToFileUrl(pdf), (prop("FilterName", "writer_pdf_Export"),))
doc.close(True)
try:
    desktop.terminate()
except Exception:
    pass
proc.wait(timeout=30)
'''


def render_pdf():
    with tempfile.NamedTemporaryFile("w", suffix=".py", delete=False) as f:
        f.write(UNO_SCRIPT)
        script = f.name
    subprocess.run(["/usr/bin/python3", script, OUT_DOCX, OUT_PDF], check=True)
    os.unlink(script)


def main():
    build_reference()
    text = "\n\n".join(open(c, encoding="utf-8").read() for c in CHAPTERS)
    md = os.path.join(HERE, "Folium_Blackbook.md")
    with open(md, "w", encoding="utf-8") as f:
        f.write(preprocess(text))
    subprocess.run(["pandoc", md, "-f", "markdown+tex_math_dollars+pipe_tables", "-t", "docx",
                    "--reference-doc", REF_DOCX, "--resource-path", HERE, "-o", OUT_DOCX], check=True, cwd=HERE)
    postprocess(OUT_DOCX)
    if "--no-pdf" not in sys.argv:
        render_pdf()
    print("built", OUT_DOCX)


if __name__ == "__main__":
    main()
